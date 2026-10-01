from docx import Document


def extract_text_from_docx(file_path):
    document = Document(file_path)

    text_parts = []

    # Extract normal paragraphs
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text)

    # Extract text from tables
    for table in document.tables:
        for row in table.rows:

            row_text = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                text_parts.append(" | ".join(row_text))

    return "\n".join(text_parts)